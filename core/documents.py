"""İzin formu çıktısı: Word (.docx) ve PDF."""
from __future__ import annotations

import os
from functools import lru_cache
from io import BytesIO
from typing import Any

from . import rules
from .utils import fmt_date, fmt_dt, now_str


def _rows(r: dict[str, Any], person: dict[str, Any]) -> list[tuple[str, str]]:
    return [
        ("Talep No", r["id"]),
        ("Personel", r["personel_adi"]),
        ("Birim / Ekip", f"{person.get('birim') or '—'} / {person.get('ekip') or '—'}"),
        ("İzin Türü", r["izin_turu"]),
        ("Başlangıç", fmt_date(r["baslangic"])),
        ("Bitiş", fmt_date(r["bitis"])),
        ("Süre türü", rules.sure_etiket(r)),
        ("Takvim Günü / İş Günü", f"{r['gun_sayisi']} / {rules.fmt_days(r['is_gunu'])}"),
        ("Açıklama", r.get("aciklama") or "—"),
        ("Durum", r["durum"]),
        ("Karar Veren / Zamanı", f"{r.get('karar_veren') or '—'} / {fmt_dt(r.get('onay_zamani'))}"),
        ("Talebi Oluşturan / Zamanı", f"{r['olusturan']} / {fmt_dt(r['olusturma_zamani'])}"),
    ]


def build_docx(r: dict[str, Any], person: dict[str, Any], verify_url: str | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading("İZİN FORMU", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in _rows(r, person):
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = str(v)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph()
    sign = doc.add_table(rows=2, cols=3)
    sign.style = "Table Grid"
    for i, h in enumerate(["Personel", "Yönetici", "İK Teslim Alan"]):
        sign.rows[0].cells[i].text = h
        sign.rows[1].cells[i].text = "\n\nAd Soyad / İmza / Tarih"
    doc.add_paragraph()
    if verify_url:
        from docx.shared import Cm
        from . import verify
        doc.add_picture(BytesIO(verify.qr_png(verify_url)), width=Cm(2.8))
        doc.add_paragraph("Belgeyi doğrulamak için QR kodu okutun: " + verify_url)
    doc.add_paragraph(f"Bu belge İzin Takip Sistemi tarafından {now_str()} tarihinde üretilmiştir. Talep No: {r['id']}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@lru_cache(maxsize=1)
def _find_font() -> str | None:
    candidates = [
        os.environ.get("IZIN_FONT_PATH"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    return next((c for c in candidates if c and os.path.exists(c)), None)


def pdf_available() -> bool:
    try:
        import fpdf  # noqa: F401
    except ImportError:
        return False
    return _find_font() is not None


def build_pdf(r: dict[str, Any], person: dict[str, Any], verify_url: str | None = None) -> bytes | None:
    """Türkçe karakterli bir TTF bulunamazsa None döner (Word çıktısı yine kullanılabilir)."""
    if not pdf_available():
        return None
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("U", "", _find_font())
    pdf.set_font("U", size=16)
    pdf.cell(0, 14, "İZİN FORMU", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("U", size=11)
    for k, v in _rows(r, person):
        pdf.cell(60, 9, k, border=1)
        pdf.multi_cell(0, 9, str(v), border=1, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(12)
    w = (pdf.w - pdf.l_margin - pdf.r_margin) / 3
    for h in ["Personel", "Yönetici", "İK Teslim Alan"]:
        pdf.cell(w, 9, h, border=1, align="C")
    pdf.ln()
    for _ in range(3):
        pdf.cell(w, 22, "Ad Soyad / İmza / Tarih", border=1, align="C")
    pdf.ln(8)
    if verify_url:
        from . import verify
        pdf.image(BytesIO(verify.qr_png(verify_url)), w=28)
        pdf.set_font("U", size=8)
        pdf.multi_cell(0, 5, "Belgeyi doğrulamak için QR kodu okutun: " + verify_url, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    pdf.set_font("U", size=9)
    pdf.cell(0, 6, f"Bu belge İzin Takip Sistemi tarafından {now_str()} tarihinde üretilmiştir. Talep No: {r['id']}",
             new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())

# END-OF-FILE v3.2
